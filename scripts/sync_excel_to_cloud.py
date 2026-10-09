#!/usr/bin/env python3
"""
scripts/sync_excel_to_cloud.py
------------------------------
A single CLI script that performs both steps:
  1. Reads an Excel file (.xlsx) and converts it to JSON format (matching master catalog schema).
  2. Upserts the JSON products directly into Supabase Cloud Database under a specific merchant shop.

Usage:
  # Convert Excel -> JSON -> Supabase Cloud (Default YesFancy shop, draft mode)
  python3 scripts/sync_excel_to_cloud.py --excel data/draft_catalog_import.xlsx

  # Only convert Excel to JSON without pushing to cloud:
  python3 scripts/sync_excel_to_cloud.py --excel data/draft_catalog_import.xlsx --dry-run

  # Target a specific shop slug:
  python3 scripts/sync_excel_to_cloud.py --excel data/draft_catalog_import.xlsx --shop yesfancy
"""

import os
import sys
import json
import ssl
import argparse
import urllib.request
import urllib.error
import openpyxl

SHOP_IDS = {
    "yesfancy": "8ca39196-2948-4d61-bfe2-388875120b2b",
    "varasiddhi": "2a9a8732-f2c0-48ee-8fc0-63918c02c8bb",
    "mumbai_paan": "f0439b68-8ed1-4880-9d14-a9a0cdbaffd8"
}

def load_env():
    env_vars = {}
    for env_file in ['.env', '.env.local', '.env.production']:
        if os.path.exists(env_file):
            with open(env_file) as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith('#') and '=' in line:
                        k, v = line.split('=', 1)
                        env_vars[k.strip()] = v.strip().strip('"').strip("'")
    return env_vars

def excel_to_json(excel_path: str, output_json_path: str):
    print(f"[*] Reading Excel workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    headers = [str(c.value or '').strip() for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]

    products = []
    for r in rows:
        sku = str(r.get('SKU ID') or '').strip()
        if not sku:
            continue
        
        # truncate SKU if > 55 chars to fit database varchar(64)
        if len(sku) > 55:
            sku = sku[:55].rstrip('_')

        tags_raw = r.get('Tags') or ''
        tags = [t.strip() for t in str(tags_raw).split(',') if t.strip()]

        cat_id = str(r.get('Category ID') or 'novelties').strip().lower()
        if not tags:
            tags = [cat_id]

        is_active = r.get('Is Active')
        if isinstance(is_active, str):
            is_active = is_active.lower() in ['true', '1', 'yes']
        elif is_active is None:
            is_active = False
        else:
            is_active = bool(is_active)

        mrp = float(r.get('MRP (₹)') or 0)
        disc_raw = r.get('Discount (%)')
        discount = float(disc_raw) if disc_raw is not None else 0

        prod = {
            "id": sku,
            "name": str(r.get('Product Title') or '').strip(),
            "brand": str(r.get('Brand') or 'Generic').strip(),
            "category_id": cat_id,
            "sub_category": str(r.get('Sub-Category') or 'General').strip(),
            "tags": tags,
            "description": str(r.get('Description') or '').strip(),
            "mrp": mrp,
            "discount": discount,
            "capacity": str(r.get('Capacity / Size') or '').strip(),
            "image_url": str(r.get('Image URL (Upload in Backend)') or '').strip(),
            "badge": str(r.get('Badge') or '').strip() or None,
            "is_active": is_active,
            "stock": int(r.get('Stock') or 50),
            "has_variants": False,
            "variants": []
        }
        products.append(prod)

    with open(output_json_path, 'w', encoding='utf-8') as f:
        json.dump(products, f, indent=2, ensure_ascii=False)

    print(f"[✓] Converted {len(products)} products from Excel -> JSON: {output_json_path}")
    return products

def json_to_supabase(products: list, shop_id: str, supabase_url: str, service_key: str):
    print(f"[*] Synchronizing {len(products)} products to Supabase Cloud (shop_id: {shop_id})...")
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE

    db_rows = []
    for p in products:
        mrp = float(p.get("mrp", 0))
        discount = float(p.get("discount", 0))
        computed_price = round(mrp * (1 - discount / 100))

        db_rows.append({
            "id": p["id"],
            "shop_id": shop_id,
            "name": p["name"],
            "description": p.get("description", ""),
            "mrp": mrp,
            "discount": discount,
            "price": computed_price,
            "category_id": p.get("category_id", "novelties"),
            "badge": p.get("badge") or None,
            "tags": p.get("tags") if isinstance(p.get("tags"), list) else [p.get("category_id", "novelties")],
            "image_url": p.get("image_url", ""),
            "stock": int(p.get("stock", 50)),
            "is_active": bool(p.get("is_active", False))
        })

    batch_size = 50
    total_upserted = 0
    url = f"{supabase_url}/rest/v1/products?on_conflict=shop_id,id"

    for i in range(0, len(db_rows), batch_size):
        chunk = db_rows[i:i + batch_size]
        req = urllib.request.Request(
            url,
            data=json.dumps(chunk).encode("utf-8"),
            headers={
                "apikey": service_key,
                "Authorization": f"Bearer {service_key}",
                "Content-Type": "application/json",
                "Prefer": "resolution=merge-duplicates"
            },
            method="POST"
        )
        try:
            with urllib.request.urlopen(req, context=ctx) as resp:
                total_upserted += len(chunk)
                print(f"  ✓ Batch {i//batch_size + 1}: Upserted {len(chunk)} items ({total_upserted}/{len(db_rows)})")
        except urllib.error.HTTPError as e:
            print(f"  ❌ Error in batch {i//batch_size + 1}: {e.code} - {e.read().decode('utf-8')}")

    print(f"[✓] Supabase synchronization complete! Total upserted: {total_upserted}")

def main():
    parser = argparse.ArgumentParser(description="Convert Excel catalog to JSON and sync to Supabase Cloud.")
    parser.add_argument("--excel", default="data/draft_catalog_import.xlsx", help="Path to input Excel file")
    parser.add_argument("--json-out", default="data/draft_catalog_import.json", help="Path to output JSON file")
    parser.add_argument("--shop", default="yesfancy", choices=list(SHOP_IDS.keys()), help="Target merchant shop")
    parser.add_argument("--dry-run", action="store_true", help="Only generate JSON, do not push to cloud DB")
    args = parser.parse_args()

    products = excel_to_json(args.excel, args.json_out)

    if args.dry_run:
        print("[!] Dry run enabled: Skipping Supabase Cloud push.")
        return

    env = load_env()
    supabase_url = env.get("PUBLIC_SUPABASE_URL", "https://fwrievuhudjkeffmszqf.supabase.co")
    service_key = env.get("SUPABASE_SERVICE_ROLE_KEY", env.get("PUBLIC_SUPABASE_ANON_KEY"))

    if not service_key:
        print("❌ Error: Supabase credentials not found in environment.")
        return

    shop_id = SHOP_IDS[args.shop]
    json_to_supabase(products, shop_id, supabase_url, service_key)

if __name__ == "__main__":
    main()
